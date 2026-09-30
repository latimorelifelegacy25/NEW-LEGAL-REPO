from __future__ import annotations

from dataclasses import dataclass, asdict
from io import BytesIO
from typing import Iterable

from docx import Document

@dataclass(frozen=True)
class DocxParagraph:
    index: int
    text: str
    style: str | None

def read_docx_paragraphs(data: bytes) -> list[dict]:
    doc=Document(BytesIO(data))
    out=[]
    for idx,p in enumerate(doc.paragraphs):
        out.append(asdict(DocxParagraph(
            index=idx,
            text=p.text,
            style=p.style.name if p.style else None,
        )))
    return out

def apply_approved_edits(
    *,
    original_data: bytes,
    approved_edits: Iterable[dict],
) -> bytes:
    doc=Document(BytesIO(original_data))
    paragraphs=doc.paragraphs
    for edit in approved_edits:
        if edit.get("status") != "approved":
            continue
        location=edit.get("location", "")
        if not location.startswith("paragraph-index:"):
            raise ValueError(f"Unsupported edit location: {location}")
        try:
            idx=int(location.split(":",1)[1])
        except ValueError as exc:
            raise ValueError(f"Invalid paragraph index: {location}") from exc
        if idx < 0 or idx >= len(paragraphs):
            raise ValueError(f"Paragraph index out of range: {idx}")
        expected=str(edit.get("before",""))
        if paragraphs[idx].text != expected:
            raise ValueError(f"Edit conflict at paragraph-index:{idx}: source text no longer matches approved 'before' text")
        paragraphs[idx].text=str(edit.get("after",""))
    out=BytesIO()
    doc.save(out)
    return out.getvalue()


def read_numbered_docx_paragraphs(data: bytes) -> list[dict]:
    """Return Word list-numbered paragraphs in document order.

    Word often stores pleading numbers in w:numPr rather than literal paragraph text.
    This detects those paragraphs without pretending the visible number is embedded
    in p.text. Sequential display numbers are assigned in document order and the
    underlying numId/ilvl are retained for audit.
    """
    doc=Document(BytesIO(data))
    out=[]
    display_number=0
    for idx,p in enumerate(doc.paragraphs):
        pPr=p._p.pPr
        numPr=pPr.numPr if pPr is not None else None
        if numPr is None:
            continue
        display_number += 1
        num_id=int(numPr.numId.val) if numPr.numId is not None else None
        ilvl=int(numPr.ilvl.val) if numPr.ilvl is not None else None
        out.append({
            "index": idx,
            "display_number": display_number,
            "text": p.text,
            "style": p.style.name if p.style else None,
            "num_id": num_id,
            "level": ilvl,
        })
    return out
