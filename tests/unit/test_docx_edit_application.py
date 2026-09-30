from io import BytesIO
import pytest
from docx import Document

from services.legal.docx_io import apply_approved_edits, read_docx_paragraphs
from services.legal import approvals

def make_docx() -> bytes:
    doc=Document()
    doc.add_paragraph("Paragraph one.")
    doc.add_paragraph("Paragraph two.")
    out=BytesIO()
    doc.save(out)
    return out.getvalue()

def test_read_docx_paragraphs():
    paras=read_docx_paragraphs(make_docx())
    assert paras[0]["text"]=="Paragraph one."
    assert paras[1]["text"]=="Paragraph two."

def test_approved_edit_requires_before_match():
    edit={"edit_id":"e1","status":"approved","location":"paragraph-index:0","before":"Paragraph one.","after":"Paragraph one corrected."}
    result=apply_approved_edits(original_data=make_docx(),approved_edits=[edit])
    assert read_docx_paragraphs(result)[0]["text"]=="Paragraph one corrected."

def test_stale_approved_edit_is_rejected():
    edit={"edit_id":"e2","status":"approved","location":"paragraph-index:0","before":"Different text.","after":"Replacement."}
    with pytest.raises(ValueError,match="source text no longer matches"):
        apply_approved_edits(original_data=make_docx(),approved_edits=[edit])

def test_server_approval_is_document_scoped():
    approvals.propose({"edit_id":"scope-test","document_id":"doc-a","location":"paragraph-index:0","before":"Paragraph one.","after":"Changed."})
    approvals.decide("scope-test",True)
    assert len(approvals.approved_for_document("doc-a",["scope-test"]))==1
    assert approvals.approved_for_document("doc-b",["scope-test"])==[]
